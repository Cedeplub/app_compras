"""Etapa 12 — o lote de preços como documento (v2/PROMPT_ETAPA_12_PRECOS_DEFINIDOS.md
§4, §5, §6.2; `sql/05_tabelas_lote_preco.sql`).

Molde exato de `app/servicos/pedido.py` (cabeçalho + itens + histórico de
status), aplicado à decisão de preço em vez de ao carrinho. `APP_DECISAO_PRECO`
e `APP_DECISAO_PRECO_HIST` continuam intocadas por este módulo — são gravadas
aqui só CHAMANDO `preco._gravar_decisao_preco_conn` (nunca duplicando a regra),
e continuam sendo o que volta para o modelo no próximo `dbt run` (CONTEXTO.md
§6 regra 10). O lote é a camada de cima: o documento que sai da diretoria em
direção a quem digita no Winthor (rotina 201).

── Decisão do usuário em 10/09/2026: NÃO existe status "Aplicado" ──────────
`CK_APP_LOTE_PRECO_STATUS` já foi ajustado no banco (conferido em
11/09/2026: `status in ('Rascunho', 'Enviado')`, sem 'Aplicado' no domínio) —
esta aplicação NUNCA grava esse valor de qualquer forma. Não há usuário
no app para quem opera a rotina 201, e por isso ninguém tem esse clique
para dar. O que existia como
status "Aplicado" no protótipo do Diretor virou
CONFERÊNCIA MEDIDA por item (§4.5): compara `pv_*_novo` desta tabela com
`COMPRAS_PEDIDO.PV_ATACADO`/`PV_VAREJO` (join, na leitura, tolerância 0,005) e
classifica cada item como aplicado/pendente/parcial/divergente (ver
`_situacao_item`: 'parcial' é a mistura sem alarme — um canal aplicado e o
outro ainda pendente; 'divergente' é reservado para o banco ter ido a um
terceiro valor de verdade). Por isso
`ORDEM_STATUS` tem só 2 valores, e a máquina de estados inteira é um passo só
(Rascunho <-> Enviado).

── PV_*_ATUAL é congelado (mesmo raciocínio do FATOR_EXIBICAO, MELHORIA A5) ──
O snapshot é lido de COMPRAS_PEDIDO só na hora em que o item ENTRA no lote
(primeira vez). Depois que o preço é aplicado no Winthor e o dbt reconstrói,
PV_ATACADO/PV_VAREJO no banco passam a valer o preço novo — se "atual" fosse
lido ao vivo depois disso, a coluna "de -> para" do lote mostraria "de R$
66,00 para R$ 66,00". Édito de item (`upsert_item`) NUNCA reescreve
`pv_*_atual` de um item que já existe no lote.

── Canal não decidido é NULO, não "preço igual ao atual" (§4.4) ────────────
`pv_atacado_novo`/`pv_varejo_novo` nulo == "este canal não foi tocado neste
item". O arquivo da rotina 201 de um canal só inclui itens com decisão
NAQUELE canal — nunca reempilha o preço atual como se fosse novo.

── COD_FAB é o identificador do arquivo da 201, não o CODIGO interno (§4.6) ──
Ver `IDENTIFICADOR_ROTINA_201` e `_cod_fab_repetidos`. A checagem de
duplicidade é contra o CATÁLOGO INTEIRO (`COMPRAS_PEDIDO`), não só os itens do
lote — o risco é colidir com um produto que não está no lote.

Nomes em snake_case (vocabulário do banco); a tradução para camelCase mora em
`app/api/contrato.py` (`lote_preco`, `item_lote_preco`, `pagina_lotes_preco`).
"""
from __future__ import annotations

import datetime as dt
import io
import threading

import oracledb
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.core import auditoria, database
from app.servicos import exportacao, preco

# ─────────────────────────────────────────────────────────────────────────────
# Máquina de estados — só 2 passos (ver docstring do módulo, seção "Decisão do
# usuário em 10/09/2026"). `CK_APP_LOTE_PRECO_STATUS` já foi ajustado no banco
# (conferido em 11/09/2026: só 'Rascunho'/'Enviado' no domínio, ver docstring
# do módulo) — esta aplicação nunca gravaria 'Aplicado' de qualquer forma.
# ─────────────────────────────────────────────────────────────────────────────
ORDEM_STATUS = ["Rascunho", "Enviado"]
STATUS_EDITAVEIS = {"Rascunho", "Enviado"}

_TOLERANCIA_APLICACAO = 0.005
"""Fonte única — usada tanto no SQL de agregação (`_AGG_LOTE`, via f-string)
quanto no cálculo de `situacao` em Python (`_status_canal`). Antes eram dois
literais `0.005` chumbados em lugares diferentes; mudar um sem o outro fazia
a lista ("N de M aplicados") divergir do detalhe do item em silêncio."""


class LoteNaoEncontrado(Exception):
    """id_lote não existe em APP_LOTE_PRECO. A rota converte para 404."""


class TransicaoInvalida(Exception):
    """Pulo de etapa ou volta além do início. 409."""


class RascunhoConflitante(Exception):
    """Voltar um lote Enviado para Rascunho violaria
    `UX_APP_LOTE_PRECO_RASCUNHO`: o mesmo `criado_por` já tem OUTRO lote em
    Rascunho aberto (a mensagem nomeia esse outro lote — ver
    `_conflito_rascunho`). Irmã de `TransicaoInvalida`, não a mesma classe: o
    significado é diferente ("existe um documento concorrente", não "pulou
    etapa") e a mensagem precisa do número do outro lote, que
    `TransicaoInvalida` não carrega. 409 — nunca junta os dois lotes
    sozinho, porque fundir dois documentos de decisão de preço sem pedir
    seria surpresa cara."""


class EdicaoNaoPermitida(Exception):
    """Editar/remover/acrescentar item fora de Rascunho/Enviado. 409."""


class ProdutoInvalido(Exception):
    """Código não existe em COMPRAS_PEDIDO, ou item sem nenhum preço
    informado (margem sozinha não produz item de lote — CHECK
    ck_app_lote_item_um do banco exige ao menos um canal decidido). 422."""


# ─────────────────────────────────────────────────────────────── leitura ───

_AGG_LOTE = f"""
      left join (
            select i.id_lote,
                   count(*) as qtd_itens,
                   sum(case when
                         (i.pv_atacado_novo is null
                            or abs(nvl(p.pv_atacado, -999999) - i.pv_atacado_novo) < {_TOLERANCIA_APLICACAO})
                         and (i.pv_varejo_novo is null
                            or abs(nvl(p.pv_varejo, -999999) - i.pv_varejo_novo) < {_TOLERANCIA_APLICACAO})
                       then 1 else 0 end) as qtd_aplicados
              from app_lote_preco_item i
              left join compras_pedido p on p.codigo = i.id_produto
             group by i.id_lote
           ) agg on agg.id_lote = l.id_lote
"""

_COLUNAS_LOTE = """
    l.id_lote, l.status, l.observacao, l.criado_em, l.criado_por,
    l.atualizado_em, l.atualizado_por,
    nvl(agg.qtd_itens, 0) as qtd_itens,
    nvl(agg.qtd_aplicados, 0) as qtd_aplicados
"""

# Etapa 13 (§3.2/§8.2) — mesmo desenho de `produto.ORDENACOES`/`pedido.ORDENACOES`.
ORDENACOES: dict[str, tuple[str, str]] = {
    "criadoEm": ("l.criado_em", "desc"),
    "id": ("l.id_lote", "desc"),
    "status": ("l.status", "asc"),
    "qtdItens": ("nvl(agg.qtd_itens, 0)", "desc"),
    "qtdAplicados": ("nvl(agg.qtd_aplicados, 0)", "desc"),
}
ORDENACOES_VALIDAS = set(ORDENACOES)


def _ordem(ordenacao: str, direcao: str | None) -> str:
    expressao, padrao = ORDENACOES.get(ordenacao, ORDENACOES["criadoEm"])
    d = direcao if direcao in ("asc", "desc") else padrao
    # Desempate fixo por id_lote, sempre — não alterna com a direção escolhida.
    return f"{expressao} {d} nulls last, l.id_lote desc"


def listar_lotes(filtros: dict, pagina: int, itens_por_pagina: int,
                  ordenacao: str = "criadoEm", direcao: str | None = None) -> tuple[list[dict], int]:
    """qtd_itens e qtd_aplicados agregados no BANCO (§4.5: "12 de 15
    aplicados"), com a tolerância de 0,005 aplicada no próprio SQL — custo de
    um join com COMPRAS_PEDIDO, zero tabela nova."""
    condicoes: list[str] = []
    binds: dict = {}

    status_lista = filtros.get("status") or []
    if status_lista:
        marcas = ", ".join(f":st{i}" for i in range(len(status_lista)))
        binds.update({f"st{i}": s for i, s in enumerate(status_lista)})
        condicoes.append(f"l.status in ({marcas})")

    onde = f"where {' and '.join(condicoes)}" if condicoes else ""

    total_linha = database.consultar_um(f"select count(*) as n from app_lote_preco l {onde}", binds)
    total = int(total_linha["n"]) if total_linha else 0

    pagina = max(1, pagina)
    ordem = _ordem(ordenacao, direcao)
    linhas = database.consultar(
        f"""
        select {_COLUNAS_LOTE}
          from app_lote_preco l
          {_AGG_LOTE}
          {onde}
         order by {ordem}
        offset :offset rows fetch next :limite rows only
        """,
        {**binds, "offset": (pagina - 1) * itens_por_pagina, "limite": itens_por_pagina},
    )
    return linhas, total


def obter_lote(id_lote: int) -> dict | None:
    """Só o cabeçalho, com os agregados — mesma consulta de `listar_lotes`
    para uma linha só, para o fragmento pós-gravação nunca divergir da lista."""
    return database.consultar_um(
        f"select {_COLUNAS_LOTE} from app_lote_preco l {_AGG_LOTE} where l.id_lote = :id_lote",
        {"id_lote": id_lote},
    )


def _normalizar_cod_fab(cod_fab: str | None) -> str | None:
    """ÚNICO ponto que decide como um COD_FAB é normalizado — vale para
    agrupar (`_cod_fab_repetidos`), para comparar (`_motivo_cod_fab`) e para
    GRAVAR no arquivo da 201 (`gerar_xlsx_201`). Antes eram dois critérios
    (SQL `trim()` para agrupar, `.strip()` Python para comparar, valor CRU
    para gravar) — a divergência deixava passar um par como 4851/4294
    ('HC000660486' vs ' HC000660486') como se não colidisse.

    Medido em 11/09/2026 contra `compras_pedido.cod_fab` (8.772 linhas): 4
    SKUs têm espaço só nas pontas, 1 tem TABULAÇÃO só nas pontas (por isso a
    normalização é feita aqui em Python com `.strip()`, e não com `trim()` no
    SQL — o `TRIM` do Oracle sem segundo argumento remove só o caractere
    espaço, não tabulação; o SKU 471 prova isso: `trim(cod_fab)` não muda seu
    tamanho). `str.strip()` do Python cobre espaço, tabulação e não-quebrável
    (nenhuma ocorrência de não-quebrável medida hoje, mas `.strip()` já trata
    se aparecer). Espaço DUPLO INTERNO (6 SKUs, ex. 'G/  JZZ530/Q2/BRA') NÃO é
    removido: é parte do código de fábrica em si (o padrão "G/  " e "/  /" se
    repete em vários fornecedores dessa família — não é ruído de digitação),
    então colapsar mudaria a identidade do código, não limparia ruído.

    Confirmado por medição: agrupar com esta normalização dá os MESMOS 32
    grupos que `group by trim(cod_fab)` no SQL (a hipótese conservadora
    citada no §4.6), contra 31 se o critério for o valor cru — a diferença é
    exatamente o par 4851/4294."""
    if cod_fab is None:
        return None
    normalizado = cod_fab.strip()
    return normalizado or None


def _cod_fab_repetidos() -> set[str]:
    """COD_FAB duplicado no CATÁLOGO INTEIRO — §4.6, ⚠ contra TODO
    COMPRAS_PEDIDO, não só os itens deste lote: o risco é o código do lote
    colidir com um produto que não está nele. Medido em 10/09/2026: 14 SKUs
    ativos, em 6 códigos ('120717.0.02' é o pior caso: 8170, 8171, 8172,
    8173). O agrupamento acontece EM PYTHON, com `_normalizar_cod_fab` — a
    mesma função usada para comparar e gravar (ver seu docstring) — porque o
    `trim()` do SQL não normaliza tabulação, e este é o critério que decide
    quem entra no arquivo da rotina 201."""
    linhas = database.consultar("select cod_fab from compras_pedido where cod_fab is not null")
    contagem: dict[str, int] = {}
    for r in linhas:
        cf = _normalizar_cod_fab(r["cod_fab"])
        if cf:
            contagem[cf] = contagem.get(cf, 0) + 1
    return {cf for cf, n in contagem.items() if n > 1}


def _motivo_cod_fab(cod_fab: str | None, repetidos: set[str]) -> str | None:
    """None == elegível para o arquivo da 201. Caso contrário, o motivo exato
    que a tela e o Excel de conferência mostram (§4.6)."""
    cf = _normalizar_cod_fab(cod_fab) or ""
    if not cf:
        return "sem código de fábrica"
    if cf in repetidos:
        return "código de fábrica repetido"
    return None


def _status_canal(atual, novo, banco) -> str | None:
    """§4.5, por canal. None == canal não decidido neste item."""
    if novo is None:
        return None
    novo = float(novo)
    if banco is not None and abs(float(banco) - novo) < _TOLERANCIA_APLICACAO:
        return "aplicado"
    if banco is None:
        return "pendente"  # ainda sem leitura do banco para este produto
    if atual is not None and abs(float(banco) - float(atual)) < _TOLERANCIA_APLICACAO:
        return "pendente"
    return "divergente"


def _situacao_item(item: dict) -> str | None:
    """Combina os dois canais decididos num status só (§4.5): só é 'aplicado'
    se TODOS os canais decididos baterem com o banco.

    'divergente' é reservado para o que o vocabulário do §4.5 realmente
    descreve: o banco foi para um TERCEIRO valor, diferente do atual E do
    novo — alarme de verdade. Atacado aplicado + varejo ainda pendente NÃO é
    isso: nenhum canal foi para lugar nenhum inesperado, um dos dois só ainda
    não foi digitado no Winthor. Rotular esse caso como 'divergente' soaria
    como o mesmo alarme de um valor estranho no banco, quando é só "faltou a
    metade" — por isso existe 'parcial': mistura de aplicado/pendente, SEM
    nenhum canal realmente divergente. `divergente` continua vencendo se
    aparecer em qualquer canal, porque aí sim é o caso grave."""
    status = [
        _status_canal(item.get("pv_atacado_atual"), item.get("pv_atacado_novo"), item.get("pv_atacado_banco")),
        _status_canal(item.get("pv_varejo_atual"), item.get("pv_varejo_novo"), item.get("pv_varejo_banco")),
    ]
    status = [s for s in status if s is not None]
    if not status:
        return None
    if any(s == "divergente" for s in status):
        return "divergente"
    if all(s == "aplicado" for s in status):
        return "aplicado"
    if all(s == "pendente" for s in status):
        return "pendente"
    return "parcial"


def _a_prazo_valido(atual_congelado, atual_banco, a_prazo, tolerancia: float = _TOLERANCIA_APLICACAO):
    """`ALT_PV_*_AP` vem de `COMPRAS_PEDIDO`, calculado pelo dbt a partir do
    `PV_ATACADO`/`PV_VAREJO` CORRENTE do banco (`atual_banco`) — nunca do
    preço CONGELADO que a tela mostra como "atual" (`atual_congelado`, MELHORIA
    A5). As duas coisas normalmente coincidem (é assim logo que o item entra
    no lote), mas divergem quando outro lote foi aplicado e incorporado por um
    `dbt run` DEPOIS que este item entrou neste lote: aí o a prazo do banco
    corresponde a um preço à vista que não é nem o "atual" nem o "novo" que a
    tela está mostrando. Mostrar esse número nessa hora sugere que ele está
    ligado a um dos dois preços da tela, e não está — apagar é mais honesto
    (medido em 11/09/2026: 2 SKUs no catálogo inteiro têm a coluna preenchida
    hoje; recalcular aqui seria fórmula duplicada, proibido neste projeto —
    CONTEXTO.md regra 10, `ALT_PV_*` é decisão humana replicada pelo dbt, não
    algo que a aplicação decida de novo)."""
    if a_prazo is None or atual_congelado is None or atual_banco is None:
        return None
    if abs(float(atual_congelado) - float(atual_banco)) >= tolerancia:
        return None
    return a_prazo


def obter_itens(id_lote: int) -> list[dict]:
    """Itens com descrição/cód. fábrica/departamento de COMPRAS_PEDIDO por
    LEFT JOIN (nunca INNER — mesmo raciocínio de `pedido.obter_itens`: o
    produto pode ter saído do catálogo depois que o lote foi enviado). Cada
    item ganha `situacao` (§4.5) e `entra_arquivo_201`/`motivo_exclusao_201`
    (§4.6), calculados aqui para a tela e os dois exportadores nunca
    divergirem."""
    linhas = database.consultar(
        """
        select i.id_lote, i.id_produto as codigo,
               i.pv_atacado_atual, i.pv_atacado_novo,
               i.pv_varejo_atual, i.pv_varejo_novo,
               i.criado_em, i.criado_por, i.atualizado_em, i.atualizado_por,
               p.descricao, p.cod_fab, p.fornecedor as departamento,
               p.pv_atacado as pv_atacado_banco, p.pv_varejo as pv_varejo_banco,
               p.alt_pv_at_ap, p.alt_pv_var_ap
          from app_lote_preco_item i
          left join compras_pedido p on p.codigo = i.id_produto
         where i.id_lote = :id_lote
         order by i.id_produto
        """,
        {"id_lote": id_lote},
    )
    repetidos = _cod_fab_repetidos()
    itens = []
    for linha in linhas:
        item = dict(linha)
        item["situacao"] = _situacao_item(item)
        motivo = _motivo_cod_fab(item.get("cod_fab"), repetidos)
        # Normaliza AQUI, no ponto único de leitura — item cru chegava até a
        # tela (`contrato.item_lote_preco`) e o Excel de conferência
        # (`_montar_xlsx_conferencia`) com espaço/tabulação invisível na
        # ponta (SKU 4294/471, ver docstring de `_normalizar_cod_fab`).
        # Exatamente o item excluído do arquivo 201 por repetição — o que
        # vai para digitação MANUAL — é o que mais precisa chegar limpo.
        item["cod_fab"] = _normalizar_cod_fab(item.get("cod_fab"))
        item["entra_arquivo_201"] = motivo is None
        item["motivo_exclusao_201"] = motivo
        # §4.4/§9.3 — ver docstring de `_a_prazo_valido`: apaga o a prazo
        # quando ele não corresponde mais ao preço congelado exibido como
        # "atual" (calculado aqui, uma vez, para a tela e o Excel nunca
        # mostrarem números diferentes).
        item["alt_pv_at_ap"] = _a_prazo_valido(
            item.get("pv_atacado_atual"), item.get("pv_atacado_banco"), item.get("alt_pv_at_ap")
        )
        item["alt_pv_var_ap"] = _a_prazo_valido(
            item.get("pv_varejo_atual"), item.get("pv_varejo_banco"), item.get("alt_pv_var_ap")
        )
        itens.append(item)
    return itens


def obter_detalhe(id_lote: int) -> dict | None:
    cabecalho = obter_lote(id_lote)
    if cabecalho is None:
        return None
    cabecalho = dict(cabecalho)
    cabecalho["itens"] = obter_itens(id_lote)
    return cabecalho


# ──────────────────────────────────────────────────── criar / acrescentar ───

def _normalizar_itens(itens: list[dict]) -> dict[int, dict]:
    """`itens`: [{"codigo", "margem_alvo"?, "margem_alvo_varejo"?,
    "alt_pv_at_av"?, "alt_pv_var_av"?}]. Dedupe por código (última ocorrência
    vence, mesmo raciocínio de `pedido.salvar_carrinho`/`preco.
    gravar_decisao_preco_lote`). Recusa item sem NENHUM preço — margem
    sozinha não produz uma linha de lote (CHECK ck_app_lote_item_um)."""
    normalizados: dict[int, dict] = {}
    for it in itens:
        codigo = int(it["codigo"])
        if it.get("alt_pv_at_av") is None and it.get("alt_pv_var_av") is None:
            raise ProdutoInvalido(
                f"Produto {codigo}: informe ao menos um preço (atacado ou"
                " varejo) — margem sozinha não produz item de lote."
            )
        normalizados[codigo] = it
    return normalizados


def _buscar_catalogo(codigos: list[int]) -> dict[int, dict]:
    marcas = ", ".join(f":c{i}" for i in range(len(codigos)))
    binds = {f"c{i}": c for i, c in enumerate(codigos)}
    linhas = database.consultar(
        f"select codigo, pv_atacado, pv_varejo from compras_pedido where codigo in ({marcas})",
        binds,
    )
    return {int(r["codigo"]): r for r in linhas}


def _achar_rascunho_do_usuario(cur, usuario_login: str) -> int | None:
    """§4.3 do prompt: "A tela reaproveita o Rascunho aberto do próprio
    usuário quando existe" — implementado NO SERVIDOR, não na tela, e é
    proposital: se a tela listasse os lotes, escolhesse o Rascunho e só
    depois fizesse o POST, duas abas abertas (ou dois cliques rápidos) veriam
    a MESMA lista, escolheriam o MESMO Rascunho inexistente e cada uma
    criaria o seu — dois arquivos de importação para o mesmo pedido,
    exatamente o que a regra existe para impedir. No servidor a garantia é
    estrutural porque a busca
    e o INSERT/UPDATE seguintes vivem na MESMA transação, com esta linha
    travada por FOR UPDATE: a segunda chamada concorrente espera a primeira
    liberar o lock e aí enxerga o lote que a primeira acabou de criar (ou
    passa a acrescentar no mesmo), nunca cria um segundo. Não depende de a
    tela lembrar de nada — vale para qualquer cliente (app, script, curl).

    "Do próprio usuário" é pelo LOGIN, o mesmo que já grava em `criado_por`
    (não há tabela de sessão aqui para casar por id).

    Trava (FOR UPDATE) TODOS os Rascunhos do usuário e devolve o mais
    recente via `linhas[0]`. Hoje `UX_APP_LOTE_PRECO_RASCUNHO` já garante no
    banco que existe NO MÁXIMO um Rascunho por `criado_por` — o `fetchall()`
    (em vez de `fetchone()`) é cinto sobre suspensório, não necessidade:
    fica como rede para o caso de o índice estar ausente/desatualizado numa
    base antiga, não porque o app ainda espere ver mais de um."""
    cur.execute(
        """
        select id_lote from app_lote_preco
         where status = 'Rascunho' and criado_por = :usuario
         order by criado_em desc, id_lote desc
         for update
        """,
        {"usuario": usuario_login},
    )
    linhas = cur.fetchall()
    return int(linhas[0][0]) if linhas else None


def _bloquear_lote(cur, id_lote: int) -> dict:
    """Lê e trava (FOR UPDATE) o cabeçalho — serializa gravações concorrentes
    no mesmo lote (mesmo raciocínio de `pedido._bloquear_pedido`). Devolve
    `criado_por` junto com `status` porque `_transicionar` precisa dele para
    checar `UX_APP_LOTE_PRECO_RASCUNHO` antes de voltar para Rascunho (ver
    `_conflito_rascunho`)."""
    cur.execute(
        "select status, criado_por from app_lote_preco where id_lote = :id for update",
        {"id": id_lote},
    )
    linha = cur.fetchone()
    if linha is None:
        raise LoteNaoEncontrado(f"Lote {id_lote} não encontrado.")
    return {"status": linha[0], "criado_por": linha[1]}


def _exigir_editavel(status_atual: str) -> None:
    if status_atual not in STATUS_EDITAVEIS:
        raise EdicaoNaoPermitida(
            f"Lote está em '{status_atual}'; só é possível editar em Rascunho ou Enviado."
        )


def _gravar_item_lote_conn(
    conn, id_lote: int, codigo: int, it: dict,
    pv_atacado_atual_banco, pv_varejo_atual_banco,
    usuario_login: str, usuario_id: int, ip: str | None, agora: dt.datetime,
) -> None:
    """Grava a decisão (APP_DECISAO_PRECO + histórico, pela função que já
    existe) E a linha do item do lote, na MESMA conexão/transação do
    chamador. Reaproveitado por `criar_lote`, `acrescentar_itens` e
    `upsert_item` — um único lugar que sabe como as duas tabelas se casam.

    `pv_*_atual_banco` só é usado no ramo INSERT do merge (item entrando pela
    primeira vez no lote): é o snapshot congelado. No ramo UPDATE (item já
    existia), `pv_*_atual` da linha NÃO é tocado — é assim que ele permanece
    congelado mesmo quando o item é editado depois (MELHORIA A5)."""
    preco._gravar_decisao_preco_conn(
        conn, codigo,
        it.get("margem_alvo"), it.get("margem_alvo_varejo"),
        it.get("alt_pv_at_av"), it.get("alt_pv_var_av"),
        usuario_login, usuario_id, ip, agora,
    )
    cur = conn.cursor()
    cur.execute(
        """
        merge into app_lote_preco_item t
        using (select :id_lote as id_lote, :codigo as id_produto from dual) s
           on (t.id_lote = s.id_lote and t.id_produto = s.id_produto)
         when matched then update set
              pv_atacado_novo = nvl(:pv_atacado_novo, t.pv_atacado_novo),
              pv_varejo_novo  = nvl(:pv_varejo_novo, t.pv_varejo_novo),
              atualizado_em = :agora, atualizado_por = :usuario
         when not matched then insert
              (id_lote, id_produto, pv_atacado_atual, pv_atacado_novo,
               pv_varejo_atual, pv_varejo_novo,
               criado_em, criado_por, atualizado_em, atualizado_por)
              values (:id_lote, :codigo, :pv_atacado_atual, :pv_atacado_novo,
                      :pv_varejo_atual, :pv_varejo_novo,
                      :agora, :usuario, :agora, :usuario)
        """,
        {
            "id_lote": id_lote, "codigo": codigo,
            "pv_atacado_atual": pv_atacado_atual_banco, "pv_atacado_novo": it.get("alt_pv_at_av"),
            "pv_varejo_atual": pv_varejo_atual_banco, "pv_varejo_novo": it.get("alt_pv_var_av"),
            "agora": agora, "usuario": usuario_login,
        },
    )
    cur.close()


class IndiceRascunhoAmbiguo(Exception):
    """Mais de um índice único em APP_LOTE_PRECO que não sustenta PK/UNIQUE
    constraint — `_nome_indice_unico_rascunho` se recusa a adivinhar qual é
    o `UX_APP_LOTE_PRECO_RASCUNHO` com `linhas[0]`. Erro explícito, não
    escolha arbitrária (achado do revisor, item 7)."""


_indice_rascunho_lock = threading.Lock()
_indice_rascunho_cache: str | None = None
"""Memoização de `_nome_indice_unico_rascunho` (achado do revisor, 3ª
passada, item 1): o nome do índice não muda em tempo de execução, então
depois de achado uma vez não há por que ir ao dicionário de novo — e ir de
novo é o problema, porque essa consulta roda de dentro do `except` de
`_transicionar`, com a conexão da transação de `_bloquear_lote` ainda presa
(`with database.transacao()` não saiu ainda), e `database.consultar` pega o
semáforo (`config.MAX_EXECUCOES_SIMULTANEAS`) de novo — aninhado. Só cacheia
o resultado POSITIVO: um `None` memoizado impediria o índice de ser
encontrado se o `oracle-dba` vier a criá-lo depois (cenário citado no
docstring de `_nome_indice_unico_rascunho` abaixo), então o caminho `None`
sempre consulta de novo. O nome, uma vez achado, é definitivo — índice não
é recriado com outro nome em produção. Lock só para a escrita do cache (dupla
checagem): `uvicorn` atende em paralelo e duas threads podem cair aqui ao
mesmo tempo na primeira vez; sem lock, as duas fariam a consulta (inofensivo,
mas a checagem dupla evita até isso na maioria das vezes) e poderiam gravar
o cache uma por cima da outra — inofensivo porque o valor é o mesmo, mas o
lock deixa a escrita determinística em vez de depender de sorte de closure
do GIL."""


def _nome_indice_unico_rascunho() -> str | None:
    """Lê no dicionário de dados — NUNCA chumbado, NUNCA por posição — o
    nome do índice único que o `oracle-dba` cria para impedir dois
    `Rascunho` do MESMO `criado_por` (`UX_APP_LOTE_PRECO_RASCUNHO`, ver
    `sql/05_tabelas_lote_preco.sql`). Se o índice ainda não existir no banco
    quando esta função rodar, devolve `None` — e `_e_violacao_indice_rascunho`
    abaixo nunca reconhece um `ORA-00001` como sendo dele, então o tratamento
    especial simplesmente não dispara (ver relatório da Etapa: NÃO TESTADO
    enquanto o índice não existir). Uma consulta rápida ao dicionário só
    acontece quando um `ORA-00001` já foi lançado — não a cada `criar_lote` —
    e, depois da primeira vez que o índice for encontrado, nunca mais (ver
    `_indice_rascunho_cache` acima).

    NÃO filtra pelo CONTEÚDO da expressão funcional
    (`USER_IND_EXPRESSIONS.COLUMN_EXPRESSION` é `LONG` — Oracle não aceita
    `LIKE`/comparação sobre `LONG` em SQL, `ORA-00932`; medido rodando
    `validar/validar_lote_preco.py`, que é como esse defeito foi achado).
    Em vez disso, exclui por CARACTERÍSTICA (índice único que sustenta
    PK ou UNIQUE constraint, via `user_constraints`) — não por nome chumbado
    (`PK_APP_LOTE_PRECO`), que era a crítica original do revisor: um
    segundo índice único que TAMBÉM sustente constraint continua excluído
    do mesmo jeito, sem citar nome nenhum. Se sobrar mais de um candidato
    (um segundo índice único solto, sem constraint, por outro motivo), é
    ambíguo de verdade — `IndiceRascunhoAmbiguo`, nunca `linhas[0]`."""
    global _indice_rascunho_cache
    if _indice_rascunho_cache is not None:
        return _indice_rascunho_cache
    with _indice_rascunho_lock:
        if _indice_rascunho_cache is not None:  # outra thread já resolveu
            return _indice_rascunho_cache
        linhas = database.consultar(
            """
            select ui.index_name
              from user_indexes ui
             where ui.table_name = 'APP_LOTE_PRECO'
               and ui.uniqueness = 'UNIQUE'
               and not exists (
                     select 1 from user_constraints uc
                      where uc.table_name = ui.table_name
                        and uc.index_name = ui.index_name
                        and uc.constraint_type in ('P', 'U')
                   )
            """
        )
        if not linhas:
            return None  # índice ainda não existe; não memoiza, tenta de novo na próxima
        if len(linhas) > 1:
            raise IndiceRascunhoAmbiguo(
                "Mais de um índice único em APP_LOTE_PRECO sem constraint por trás: "
                + ", ".join(l["index_name"] for l in linhas)
            )
        _indice_rascunho_cache = linhas[0]["index_name"]
        return _indice_rascunho_cache


def _e_violacao_indice_rascunho(exc: Exception) -> bool:
    """True só quando o `ORA-00001` é do índice `UX_APP_LOTE_PRECO_RASCUNHO`
    especificamente, identificado PELO NOME — nunca um `ORA-00001` genérico,
    porque outras unicidades existem no banco e engolir todas esconderia
    defeito de verdade. Usado em dois pontos: a corrida do PRIMEIRO lote de
    um usuário (`criar_lote`) e a rede de segurança de `_transicionar` ao
    voltar para Rascunho (a checagem explícita de `_conflito_rascunho` cobre
    o caso comum; isto cobre a corrida entre a checagem e o UPDATE)."""
    texto = str(exc)
    if "ORA-00001" not in texto:
        return False
    nome_indice = _nome_indice_unico_rascunho()
    return bool(nome_indice) and nome_indice.upper() in texto.upper()


def _conflito_rascunho(cur, id_lote_atual: int, criado_por: str | None) -> int | None:
    """Checagem EXPLÍCITA, antes de gravar (não só captura do erro depois):
    o mesmo `criado_por` do lote que está voltando para Rascunho já tem
    OUTRO lote em Rascunho aberto? Se sim, o UPDATE abaixo violaria
    `UX_APP_LOTE_PRECO_RASCUNHO`. `FOR UPDATE` (mesma trava de
    `_achar_rascunho_do_usuario`) serializa contra outra transação tentando
    o mesmo agora; a captura de `ORA-00001` em `_transicionar` continua
    como rede para a corrida residual (nenhum Rascunho concorrente na hora
    desta consulta, um aparece entre esta consulta e o UPDATE)."""
    cur.execute(
        """
        select id_lote from app_lote_preco
         where status = 'Rascunho' and criado_por = :criado_por and id_lote != :id_lote_atual
         order by criado_em desc, id_lote desc
         for update
        """,
        {"criado_por": criado_por, "id_lote_atual": id_lote_atual},
    )
    linha = cur.fetchone()
    return int(linha[0]) if linha else None


def _tentar_criar_ou_acrescentar(
    normalizados: dict[int, dict], codigos: list[int], por_produto: dict[int, dict],
    observacao: str | None, usuario_login: str, usuario_id: int, ip: str | None,
) -> tuple[int, bool]:
    """Uma tentativa da transação de `criar_lote`. Devolve (id_lote,
    reaproveitado). Pode levantar `oracledb.Error` com `ORA-00001` do índice
    único de Rascunho — é `criar_lote` quem decide se isso é a corrida do
    primeiro lote (e reage) ou um erro de verdade (e propaga)."""
    agora = dt.datetime.now()
    with database.transacao() as conn:
        cur = conn.cursor()
        # Busca + trava (FOR UPDATE) do Rascunho do usuário, e o INSERT do
        # lote novo quando não há um, na MESMA transação/conexão que vai
        # inserir os itens logo abaixo — é essa unidade que fecha a corrida
        # das duas abas (ver docstring de `_achar_rascunho_do_usuario`).
        # Isso NÃO fecha a corrida quando é o PRIMEIRO Rascunho do usuário
        # (zero linhas para o FOR UPDATE travar) — só o índice único (ver
        # `_e_violacao_indice_rascunho`) fecha esse caso; o INSERT abaixo é onde
        # ele, se existir, dispara o ORA-00001.
        id_lote_existente = _achar_rascunho_do_usuario(cur, usuario_login)
        reaproveitado = id_lote_existente is not None
        if reaproveitado:
            id_lote = id_lote_existente
        else:
            id_var = cur.var(int)
            cur.execute(
                """
                insert into app_lote_preco
                    (status, observacao, criado_em, criado_por, atualizado_em, atualizado_por)
                values ('Rascunho', :observacao, :agora, :usuario, :agora, :usuario)
                returning id_lote into :id
                """,
                {"observacao": observacao, "agora": agora, "usuario": usuario_login, "id": id_var},
            )
            id_lote = int(id_var.getvalue()[0])
            cur.execute(
                """
                insert into app_lote_preco_status_hist
                    (id_lote, status_anterior, status_novo, alterado_em, alterado_por)
                values (:id_lote, null, 'Rascunho', :agora, :usuario)
                """,
                {"id_lote": id_lote, "agora": agora, "usuario": usuario_login},
            )
        cur.close()

        for codigo in codigos:
            it = normalizados[codigo]
            prod = por_produto[codigo]
            _gravar_item_lote_conn(
                conn, id_lote, codigo, it, prod["pv_atacado"], prod["pv_varejo"],
                usuario_login, usuario_id, ip, agora,
            )

        if reaproveitado:
            cur2 = conn.cursor()
            # `observacao` NÃO é descartado quando o Rascunho é reaproveitado:
            # antes esta gravação simplesmente ignorava o parâmetro. `nvl`
            # preserva a observação existente quando o chamador não manda uma
            # nova (`None`) — mesmo raciocínio de "campo não enviado = não
            # mexi nisto" já usado em `upsert_item`.
            cur2.execute(
                "update app_lote_preco set atualizado_em = :agora, atualizado_por = :usuario,"
                " observacao = nvl(:observacao, observacao)"
                " where id_lote = :id",
                {"agora": agora, "usuario": usuario_login, "id": id_lote, "observacao": observacao},
            )
            cur2.close()

        auditoria.registrar(
            conn, usuario_id,
            "ACRESCENTAR_ITENS_LOTE_PRECO" if reaproveitado else "CRIAR_LOTE_PRECO",
            "APP_LOTE_PRECO", str(id_lote),
            {
                "qtd_itens": len(codigos),
                "itens": [{"codigo": c, **normalizados[c]} for c in codigos],
                "usuario": usuario_login,
                "reaproveitado": reaproveitado,
                "observacao": observacao,
            },
            ip,
        )
    return id_lote, reaproveitado


def criar_lote(
    itens: list[dict], observacao: str | None,
    usuario_login: str, usuario_id: int, ip: str | None = None,
) -> dict:
    """§4.2: decidir o preço e criar o documento são UM ato só. Uma
    transação: grava APP_DECISAO_PRECO (com histórico), garante o cabeçalho
    em Rascunho e insere os itens com o snapshot pv_atacado/pv_varejo de
    COMPRAS_PEDIDO. Ou tudo, ou nada.

    §4.3 (MELHORIA — regra que faltou no briefing original da Etapa 12):
    "criar" aqui é condicional. Sem `id_lote` explícito nesta função (ela só
    é chamada quando o chamador NÃO tem um lote em mãos — ver
    `_achar_rascunho_do_usuario` para o motivo de a busca morar no servidor,
    não na tela), primeiro procura um Rascunho do MESMO usuário (por
    `criado_por`); se existir, ACRESCENTA os itens a ele em vez de abrir
    documento novo. `dict["lote_reaproveitado"]` no retorno diz qual dos dois
    aconteceu, para a tela distinguir "criado o lote #7" de "acrescentado ao
    lote #7".

    Recusa o lote INTEIRO se algum código não existir em COMPRAS_PEDIDO,
    nomeando os SKUs — mesmo critério de `preco.gravar_decisao_preco_lote`,
    checado ANTES de abrir a transação.

    A corrida do PRIMEIRO lote de um usuário (zero Rascunhos: nada para o
    `FOR UPDATE` de `_achar_rascunho_do_usuario` travar, então duas
    requisições simultâneas passam pela busca juntas e as duas tentam
    INSERT) é fechada pelo índice único que o `oracle-dba` cria em
    `criado_por` para `status = 'Rascunho'`, identificado em runtime por
    `_nome_indice_unico_rascunho` — nunca chumbado. Quando o INSERT perde a
    corrida, a reação é ACRESCENTAR ao lote que venceu (reconsultar e
    seguir), não devolver erro: quem clicou não tem culpa de ter clicado
    junto com outra aba, e o resultado certo é um documento só."""
    normalizados = _normalizar_itens(itens)
    codigos = list(normalizados.keys())
    if not codigos:
        raise ProdutoInvalido("Nenhum item para gravar.")
    if len(codigos) > preco.LIMITE_LOTE_PRECO:
        raise ValueError(
            f"Lote com {len(codigos)} itens excede o limite de {preco.LIMITE_LOTE_PRECO}."
            " Divida em requisições menores, todas para o MESMO idLote (§4.3) —"
            " o primeiro cria, os seguintes acrescentam."
        )

    por_produto = _buscar_catalogo(codigos)
    faltando = [c for c in codigos if c not in por_produto]
    if faltando:
        raise ProdutoInvalido(
            "Produto(s) não encontrado(s) em COMPRAS_PEDIDO: "
            + ", ".join(str(c) for c in faltando)
        )

    try:
        id_lote, reaproveitado = _tentar_criar_ou_acrescentar(
            normalizados, codigos, por_produto, observacao, usuario_login, usuario_id, ip,
        )
    except oracledb.Error as exc:
        if not _e_violacao_indice_rascunho(exc):
            raise
        # Perdemos a corrida do primeiro lote: a outra requisição já
        # committou o INSERT que criou o Rascunho. Uma segunda tentativa
        # agora ENXERGA esse Rascunho em `_achar_rascunho_do_usuario` (ele já
        # existe e está committado) e acrescenta os itens nele — nunca cria
        # um segundo documento.
        id_lote, reaproveitado = _tentar_criar_ou_acrescentar(
            normalizados, codigos, por_produto, observacao, usuario_login, usuario_id, ip,
        )

    detalhe = obter_detalhe(id_lote)
    detalhe["lote_reaproveitado"] = reaproveitado
    return detalhe


def acrescentar_itens(
    id_lote: int, itens: list[dict],
    usuario_login: str, usuario_id: int, ip: str | None = None,
) -> dict:
    """§4.3: acrescenta itens a um lote EM RASCUNHO — mesma transação, mesmo
    critério de atomicidade de `criar_lote`. É o caminho que o front usa para
    partir uma gravação de mais de `preco.LIMITE_LOTE_PRECO` itens em vários
    pedidos HTTP para o MESMO documento."""
    normalizados = _normalizar_itens(itens)
    codigos = list(normalizados.keys())
    if not codigos:
        raise ProdutoInvalido("Nenhum item para acrescentar.")
    if len(codigos) > preco.LIMITE_LOTE_PRECO:
        raise ValueError(
            f"Lote com {len(codigos)} itens excede o limite de {preco.LIMITE_LOTE_PRECO}."
            " Divida em requisições menores."
        )

    por_produto = _buscar_catalogo(codigos)
    faltando = [c for c in codigos if c not in por_produto]
    if faltando:
        raise ProdutoInvalido(
            "Produto(s) não encontrado(s) em COMPRAS_PEDIDO: "
            + ", ".join(str(c) for c in faltando)
        )

    agora = dt.datetime.now()
    with database.transacao() as conn:
        cur = conn.cursor()
        info = _bloquear_lote(cur, id_lote)
        if info["status"] != "Rascunho":
            raise EdicaoNaoPermitida(
                f"Lote está em '{info['status']}'; só é possível acrescentar itens"
                " a um lote em Rascunho."
            )
        cur.close()

        for codigo in codigos:
            it = normalizados[codigo]
            prod = por_produto[codigo]
            _gravar_item_lote_conn(
                conn, id_lote, codigo, it, prod["pv_atacado"], prod["pv_varejo"],
                usuario_login, usuario_id, ip, agora,
            )

        cur2 = conn.cursor()
        cur2.execute(
            "update app_lote_preco set atualizado_em = :agora, atualizado_por = :usuario"
            " where id_lote = :id",
            {"agora": agora, "usuario": usuario_login, "id": id_lote},
        )
        cur2.close()
        auditoria.registrar(
            conn, usuario_id, "ACRESCENTAR_ITENS_LOTE_PRECO", "APP_LOTE_PRECO", str(id_lote),
            {"qtd_itens": len(codigos), "codigos": codigos, "usuario": usuario_login}, ip,
        )

    return obter_detalhe(id_lote)


# ─────────────────────────────────────────────────────── edição de item ───

def upsert_item(
    id_lote: int, codigo: int,
    margem_alvo: float | None, margem_alvo_varejo: float | None,
    alt_pv_at_av: float | None, alt_pv_var_av: float | None,
    usuario_login: str, usuario_id: int, ip: str | None = None,
) -> dict:
    """Edita (ou acrescenta) UM item — permitido em Rascunho e Enviado
    (§6.2). Regrava APP_DECISAO_PRECO: documento e modelo não podem divergir.
    Campo não enviado (`None`) significa "não mexi nisto", nunca "apague" —
    mesmo raciocínio de `preco._gravar_decisao_preco_conn`."""
    agora = dt.datetime.now()
    with database.transacao() as conn:
        cur = conn.cursor()
        info = _bloquear_lote(cur, id_lote)
        _exigir_editavel(info["status"])

        cur.execute(
            "select pv_atacado, pv_varejo from compras_pedido where codigo = :codigo",
            {"codigo": codigo},
        )
        prod = cur.fetchone()
        if prod is None:
            raise ProdutoInvalido(f"Produto {codigo} não encontrado em COMPRAS_PEDIDO.")
        pv_atacado_banco, pv_varejo_banco = prod

        cur.execute(
            "select pv_atacado_atual, pv_atacado_novo, pv_varejo_atual, pv_varejo_novo"
            " from app_lote_preco_item where id_lote = :id_lote and id_produto = :codigo",
            {"id_lote": id_lote, "codigo": codigo},
        )
        existente = cur.fetchone()
        cur.close()
        if existente is not None:
            atual_at, novo_at_old, atual_var, novo_var_old = existente
        else:
            atual_at, novo_at_old, atual_var, novo_var_old = pv_atacado_banco, None, pv_varejo_banco, None

        novo_at = alt_pv_at_av if alt_pv_at_av is not None else novo_at_old
        novo_var = alt_pv_var_av if alt_pv_var_av is not None else novo_var_old
        if novo_at is None and novo_var is None:
            raise ProdutoInvalido(
                f"Produto {codigo}: informe ao menos um preço (atacado ou varejo)."
            )

        it = {
            "margem_alvo": margem_alvo, "margem_alvo_varejo": margem_alvo_varejo,
            "alt_pv_at_av": alt_pv_at_av, "alt_pv_var_av": alt_pv_var_av,
        }
        _gravar_item_lote_conn(
            conn, id_lote, codigo, it, atual_at, atual_var,
            usuario_login, usuario_id, ip, agora,
        )

        cur3 = conn.cursor()
        cur3.execute(
            "update app_lote_preco set atualizado_em = :agora, atualizado_por = :usuario"
            " where id_lote = :id",
            {"agora": agora, "usuario": usuario_login, "id": id_lote},
        )
        cur3.close()
        auditoria.registrar(
            conn, usuario_id, "EDITAR_ITEM_LOTE_PRECO", "APP_LOTE_PRECO_ITEM", f"{id_lote}/{codigo}",
            {"alt_pv_at_av": alt_pv_at_av, "alt_pv_var_av": alt_pv_var_av, "usuario": usuario_login}, ip,
        )

    return obter_detalhe(id_lote)


def remover_item(
    id_lote: int, codigo: int, usuario_login: str, usuario_id: int, ip: str | None = None,
) -> dict:
    """Remove o item do LOTE (o documento) — NÃO desfaz a decisão em
    APP_DECISAO_PRECO: o preço decidido permanece gravado, com histórico. O
    item só sai do arquivo/planilha de importação no Winthor."""
    agora = dt.datetime.now()
    with database.transacao() as conn:
        cur = conn.cursor()
        info = _bloquear_lote(cur, id_lote)
        _exigir_editavel(info["status"])
        cur.execute(
            "delete from app_lote_preco_item where id_lote = :id and id_produto = :codigo",
            {"id": id_lote, "codigo": codigo},
        )
        removeu = cur.rowcount > 0
        cur.execute(
            "update app_lote_preco set atualizado_em = :agora, atualizado_por = :usuario"
            " where id_lote = :id",
            {"agora": agora, "usuario": usuario_login, "id": id_lote},
        )
        cur.close()
        if removeu:
            auditoria.registrar(
                conn, usuario_id, "REMOVER_ITEM_LOTE_PRECO", "APP_LOTE_PRECO_ITEM", f"{id_lote}/{codigo}",
                {"usuario": usuario_login}, ip,
            )

    detalhe = obter_detalhe(id_lote)
    if detalhe is None:
        raise LoteNaoEncontrado(f"Lote {id_lote} não encontrado.")
    return detalhe


# ────────────────────────────────────────────────────── máquina de estados ─

def _transicionar(
    id_lote: int, passo: int, acao: str, usuario_login: str, usuario_id: int, ip: str | None
) -> dict:
    agora = dt.datetime.now()
    with database.transacao() as conn:
        cur = conn.cursor()
        info = _bloquear_lote(cur, id_lote)
        idx = ORDEM_STATUS.index(info["status"])
        novo_idx = idx + passo
        if novo_idx < 0 or novo_idx >= len(ORDEM_STATUS):
            if passo < 0:
                raise TransicaoInvalida(
                    f"Lote já está em '{info['status']}' (etapa inicial); não há como desfazer."
                )
            raise TransicaoInvalida(
                f"Lote já está em '{info['status']}' (última etapa); não há próxima etapa."
            )
        novo_status = ORDEM_STATUS[novo_idx]

        # Voltar para Rascunho está sujeito a UX_APP_LOTE_PRECO_RASCUNHO (o
        # mesmo índice que `criar_lote` já respeita) — checagem EXPLÍCITA
        # ANTES de gravar, na MESMA transação/trava de `_conflito_rascunho`,
        # para dar mensagem acionável (nomeando o outro lote) em vez de
        # deixar o ORA-00001 subir cru como 500 (ver
        # `RascunhoConflitante`). Recusa; nunca funde os dois lotes sozinho.
        if novo_status == "Rascunho":
            outro_id_lote = _conflito_rascunho(cur, id_lote, info["criado_por"])
            if outro_id_lote is not None:
                raise RascunhoConflitante(
                    f"Você já tem o lote #{outro_id_lote} em Rascunho."
                    f" Envie-o ou exclua-o antes de reabrir o #{id_lote}."
                )

        try:
            cur.execute(
                "update app_lote_preco set status = :novo, atualizado_em = :agora,"
                " atualizado_por = :usuario where id_lote = :id",
                {"novo": novo_status, "agora": agora, "usuario": usuario_login, "id": id_lote},
            )
            cur.execute(
                """
                insert into app_lote_preco_status_hist
                    (id_lote, status_anterior, status_novo, alterado_em, alterado_por)
                values (:id, :anterior, :novo, :agora, :usuario)
                """,
                {"id": id_lote, "anterior": info["status"], "novo": novo_status, "agora": agora, "usuario": usuario_login},
            )
        except oracledb.Error as exc:
            # Rede, não a checagem principal (que já rodou acima): cobre só
            # a corrida residual entre a consulta de `_conflito_rascunho` e
            # este UPDATE (outra transação cria um Rascunho concorrente
            # exatamente nesse intervalo). Sem o nome do outro lote — a
            # checagem explícita é quem dá a mensagem completa; isto só
            # evita o 500 cru se a corrida acontecer.
            if novo_status == "Rascunho" and _e_violacao_indice_rascunho(exc):
                raise RascunhoConflitante(
                    "Já existe outro lote em Rascunho para este usuário"
                    " (detectado no instante da gravação); atualize a lista"
                    " e tente novamente."
                ) from exc
            raise
        cur.close()
        auditoria.registrar(
            conn, usuario_id, acao, "APP_LOTE_PRECO", str(id_lote),
            {"de": info["status"], "para": novo_status, "usuario": usuario_login}, ip,
        )
    return obter_detalhe(id_lote)


def avancar_status(id_lote: int, usuario_login: str, usuario_id: int, ip: str | None = None) -> dict:
    return _transicionar(id_lote, +1, "AVANCAR_STATUS_LOTE_PRECO", usuario_login, usuario_id, ip)


def voltar_status(id_lote: int, usuario_login: str, usuario_id: int, ip: str | None = None) -> dict:
    return _transicionar(id_lote, -1, "VOLTAR_STATUS_LOTE_PRECO", usuario_login, usuario_id, ip)


def excluir_lote(id_lote: int, usuario_login: str, usuario_id: int, ip: str | None = None) -> None:
    """Apaga o DOCUMENTO (ON DELETE CASCADE derruba itens e histórico) —
    NUNCA a decisão em APP_DECISAO_PRECO, que continua valendo e indo para o
    próximo dbt run."""
    with database.transacao() as conn:
        cur = conn.cursor()
        info = _bloquear_lote(cur, id_lote)
        cur.execute("delete from app_lote_preco where id_lote = :id", {"id": id_lote})
        cur.close()
        auditoria.registrar(
            conn, usuario_id, "EXCLUIR_LOTE_PRECO", "APP_LOTE_PRECO", str(id_lote),
            {"status": info["status"], "usuario": usuario_login}, ip,
        )


# ──────────────────────────────────────────────────────────── exportações ──

_COLUNAS_CONFERENCIA = [
    ("CÓDIGO CEDEP", "0"),
    ("CÓD. FÁBRICA", "@"),
    ("PRODUTO", "@"),
    ("DEPARTAMENTO", "@"),
    ("ATACADO ATUAL", "#,##0.00"),
    ("ATACADO NOVO", "#,##0.00"),
    ("Δ% ATACADO", "#,##0.00"),
    ("VAREJO ATUAL", "#,##0.00"),
    ("VAREJO NOVO", "#,##0.00"),
    ("Δ% VAREJO", "#,##0.00"),
    ("A PRAZO ATACADO (INFORMATIVO, ÚLTIMO BUILD)", "#,##0.00"),
    ("A PRAZO VAREJO (INFORMATIVO, ÚLTIMO BUILD)", "#,##0.00"),
    ("ENTRA NO ARQUIVO 201?", "@"),
]


def _f_ou_none(valor):
    return None if valor is None else float(valor)


def _delta_pct(atual, novo):
    if atual is None or novo is None or float(atual) == 0:
        return None
    return (float(novo) - float(atual)) / float(atual) * 100


def _montar_xlsx_conferencia(cabecalho: dict, itens: list[dict]) -> bytes:
    """Excel para conferência e, quando necessário, digitação manual do que
    o arquivo da 201 não carrega (§4.6) — mesmo estilo visual de
    `exportacao.py` (_HEADER_FILL/_HEADER_FONT)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Conferência"

    linhas_cabecalho = [
        (f"LOTE DE PREÇOS #{cabecalho['id_lote']} — {cabecalho['status']}", True),
        (
            f"Criado em {cabecalho['criado_em']:%d/%m/%Y %H:%M} por {cabecalho.get('criado_por') or '-'}",
            False,
        ),
    ]
    if cabecalho.get("observacao"):
        linhas_cabecalho.append((f"Observação: {cabecalho['observacao']}", False))
    linhas_cabecalho.append((
        "A coluna A PRAZO é INFORMATIVA e reflete o ÚLTIMO BUILD do modelo, não"
        " o preço à vista deste lote — quem calcula o valor a prazo de verdade"
        " é o Winthor, a partir do preço à vista importado (rotina 201). Sai em"
        " branco quando o preço atual congelado já não bate com o do banco (a"
        " coluna deixaria de corresponder a este lote).",
        False,
    ))
    for texto, destaque in linhas_cabecalho:
        ws.append([texto])
        cel = ws.cell(row=ws.max_row, column=1)
        cel.font = Font(bold=destaque, italic=not destaque, size=13 if destaque else 10)
    ws.append([])

    linha_titulo = ws.max_row + 1
    ws.append([rotulo for rotulo, _ in _COLUNAS_CONFERENCIA])
    for idx in range(1, len(_COLUNAS_CONFERENCIA) + 1):
        cel = ws.cell(row=linha_titulo, column=idx)
        cel.fill = exportacao._HEADER_FILL
        cel.font = exportacao._HEADER_FONT
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    larguras = [max(len(rotulo) + 2, 12) for rotulo, _ in _COLUNAS_CONFERENCIA]

    for item in itens:
        motivo = item.get("motivo_exclusao_201")
        entra = "sim" if motivo is None else motivo
        valores = [
            item["codigo"], item.get("cod_fab"), item.get("descricao"), item.get("departamento"),
            _f_ou_none(item.get("pv_atacado_atual")), _f_ou_none(item.get("pv_atacado_novo")),
            _delta_pct(item.get("pv_atacado_atual"), item.get("pv_atacado_novo")),
            _f_ou_none(item.get("pv_varejo_atual")), _f_ou_none(item.get("pv_varejo_novo")),
            _delta_pct(item.get("pv_varejo_atual"), item.get("pv_varejo_novo")),
            _f_ou_none(item.get("alt_pv_at_ap")), _f_ou_none(item.get("alt_pv_var_ap")),
            entra,
        ]
        ws.append(valores)
        for i, valor in enumerate(valores):
            if valor is not None:
                larguras[i] = min(max(larguras[i], len(str(valor)) + 2), 45)

    ultima = ws.max_row
    for idx, (_, fmt) in enumerate(_COLUNAS_CONFERENCIA, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = larguras[idx - 1]
        if fmt == "@":
            continue
        for row in range(linha_titulo + 1, ultima + 1):
            ws.cell(row=row, column=idx).number_format = fmt

    ws.freeze_panes = f"A{linha_titulo + 1}"
    ws.auto_filter.ref = f"A{linha_titulo}:{get_column_letter(len(_COLUNAS_CONFERENCIA))}{max(ultima, linha_titulo)}"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def gerar_excel_conferencia(id_lote: int) -> tuple[bytes, str] | None:
    cabecalho = obter_lote(id_lote)
    if cabecalho is None:
        return None
    itens = obter_itens(id_lote)
    conteudo = _montar_xlsx_conferencia(cabecalho, itens)
    nome = f"lote_preco_{id_lote}_conferencia_{dt.date.today():%Y-%m-%d}.xlsx"
    return conteudo, nome


# Rotina 201 do Winthor pede, na coluna A, um entre 4 identificadores
# possíveis (documentação TOTVS, confirmado com o Diretor em 10/09/2026,
# §4.6 do prompt da Etapa 12):
#   1 - Código de Barras (UN - Venda)   -> PCEMBALAGEM, NÃO é source do dbt
#   2 - Código de Barras (UN - Master)  -> PCEMBALAGEM, NÃO é source do dbt
#   3 - Cód. de Fab. (Rotina 203)       -> COMPRAS_PEDIDO.COD_FAB  <- É ESTA
#   4 - Cód. de Fab. (Rotina 253)       -> não está no modelo
# Se um dia a importação da rotina 201 passar a ser por código de barras, a
# troca é NESTA constante — e antes disso é GRANT + stg_embalagem + coluna
# nova em COMPRAS_PEDIDO (dbt), nunca um SELECT direto ao CEDEP daqui.
IDENTIFICADOR_ROTINA_201 = "cod_fab"

_REGIAO_CANAL = {"atacado": "regiao2", "varejo": "regiao1"}


def gerar_xlsx_201(
    id_lote: int, canal: str, usuario_login: str, usuario_id: int, ip: str | None = None,
) -> tuple[bytes, str, list[dict]] | None:
    """Arquivo de importação da rotina 201: EXATAMENTE 2 colunas, SEM
    CABEÇALHO — coluna A o COD_FAB (nunca o CODIGO interno, §4.6), coluna B o
    preço à vista com 2 casas. Só entram itens com decisão NAQUELE canal
    (§4.4) e COD_FAB elegível (preenchido e não repetido no catálogo inteiro,
    §4.6). Devolve também a lista de excluídos com o motivo, para a
    auditoria (a tela já tem essa mesma informação por item via
    `obter_itens`/`obter_detalhe`, para não depender de abrir o arquivo)."""
    if canal not in _REGIAO_CANAL:
        raise ValueError(f"Canal inválido: {canal!r} (use 'atacado' ou 'varejo').")

    cabecalho = obter_lote(id_lote)
    if cabecalho is None:
        return None
    itens = obter_itens(id_lote)
    campo_novo = "pv_atacado_novo" if canal == "atacado" else "pv_varejo_novo"

    incluidos: list[tuple[str, float]] = []
    excluidos: list[dict] = []
    for item in itens:
        preco_novo = item.get(campo_novo)
        if preco_novo is None:
            continue  # canal não decidido neste item (§4.4) — nem entra, nem é "excluído"
        motivo = item.get("motivo_exclusao_201")
        if motivo is not None:
            excluidos.append({"codigo": item["codigo"], "codFab": item.get("cod_fab"), "motivo": motivo})
            continue
        # `item["cod_fab"]` já sai NORMALIZADO de `obter_itens` (mesma função
        # que decidiu que este item é elegível) — nunca o valor cru: se a
        # rotina 201 não faz trim ao casar (hipótese conservadora, ver
        # docstring de `_normalizar_cod_fab`), um espaço na ponta faz o
        # preço simplesmente não entrar, em silêncio.
        incluidos.append((item["cod_fab"], round(float(preco_novo), 2)))

    wb = Workbook()
    ws = wb.active
    for cod_fab, preco_valor in incluidos:
        ws.append([cod_fab, preco_valor])
    for row in range(1, ws.max_row + 1):
        ws.cell(row=row, column=2).number_format = "0.00"
    buffer = io.BytesIO()
    wb.save(buffer)
    conteudo = buffer.getvalue()

    nome = f"winthor_201_{canal}_{_REGIAO_CANAL[canal]}_codfab203_{id_lote}.xlsx"

    # §4.7: cada download vira linha de auditoria com o canal — baixar não
    # avança o status (são dois arquivos; ver docstring de `pedido.
    # exportar_winthor` para o contraste com o precedente do pedido).
    auditoria.registrar_isolado(
        usuario_id, "EXPORTAR_201_LOTE_PRECO", "APP_LOTE_PRECO", str(id_lote),
        {
            "canal": canal, "qtd_incluidos": len(incluidos), "qtd_excluidos": len(excluidos),
            "excluidos": excluidos, "usuario": usuario_login,
        },
        ip,
    )
    return conteudo, nome, excluidos
